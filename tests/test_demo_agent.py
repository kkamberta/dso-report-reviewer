from pathlib import Path
import sys

import openpyxl

sys.path.insert(0, str(Path(__file__).parent.parent))

from demo_agent import apply_review_decision, create_demo_job, demo_findings, reset_demo


def test_demo_findings_include_human_gate_cases():
    findings = demo_findings("ap2065", "high")

    assert len(findings) == 4
    assert any(f["flag"] == "Severity changed" and f["review_status"] == "pending" for f in findings)
    assert any(f["flag"] == "Validation failed" for f in findings)
    assert all(f["asset_importance"] == "High" for f in findings)
    row_42 = next(f for f in findings if f["row_id"] == 42)
    assert row_42["remediation_commands"][0]["label"] == "Maven dependency pin"
    assert "tomcat-embed-core" in row_42["remediation_commands"][0]["command"]


def test_create_demo_job_writes_outputs(tmp_path: Path):
    job = create_demo_job(
        {
            "job_id": "build-demo",
            "project": "ap2065",
            "priority": "high",
            "scanner": "mixed",
            "artifact_url": "https://ci.example/artifacts/report.xlsx",
        },
        tmp_path,
    )

    assert job.status == "completed"
    assert job.human_review_required is True
    assert job.validation_failed_rows == 1
    assert job.severity_review_rows == 1
    assert Path(job.output_md).exists()
    assert Path(job.output_xlsx).exists()

    summary = Path(job.output_md).read_text(encoding="utf-8")
    assert "DSO Report Reviewer Demo Summary" in summary
    assert "Human Review Queue" in summary
    assert "Technical command" in summary
    assert "mvn versions:use-dep-version" in summary

    wb = openpyxl.load_workbook(job.output_xlsx)
    ws = wb["Reviewed Findings"]
    assert ws.max_row == 5
    assert ws.cell(row=1, column=1).value == "Scanner"
    assert ws.cell(row=1, column=9).value == "Technical Commands"
    assert ws.cell(row=1, column=12).value == "Reviewer Note"
    row_42 = next(row for row in ws.iter_rows(values_only=True) if row[1] == 42)
    assert "Maven dependency pin" in row_42[8]


def test_apply_review_decision_rewrites_download_outputs(tmp_path: Path):
    job = create_demo_job(
        {
            "job_id": "build-demo",
            "project": "ap2065",
            "priority": "high",
            "scanner": "mixed",
            "artifact_url": "https://ci.example/artifacts/report.xlsx",
        },
        tmp_path,
    )

    result = apply_review_decision(
        job,
        {
            "row_id": 18,
            "review_status": "approved",
            "reviewer_note": "Approved after confirming credential rotation plan.",
        },
    )

    assert not isinstance(result, str)
    assert result["review_status"] == "approved"
    assert result["reviewer_note"] == "Approved after confirming credential rotation plan."

    summary = Path(job.output_md).read_text(encoding="utf-8")
    assert "Approved Review Decisions" in summary
    assert "Approved after confirming credential rotation plan." in summary
    assert "Row 18: Severity changed" not in summary

    wb = openpyxl.load_workbook(job.output_xlsx)
    ws = wb["Reviewed Findings"]
    row_18 = next(row for row in ws.iter_rows(values_only=True) if row[1] == 18)
    assert row_18[10] == "approved"
    assert row_18[11] == "Approved after confirming credential rotation plan."


def test_reset_demo_restores_clean_pre_approval_outputs(tmp_path: Path):
    job = reset_demo(tmp_path)
    apply_review_decision(
        job,
        {
            "row_id": 18,
            "review_status": "approved",
            "reviewer_note": "Approved after confirming credential rotation plan.",
        },
    )

    approved_summary = Path(job.output_md).read_text(encoding="utf-8")
    assert "Row 18: approved" in approved_summary

    reset_job = reset_demo(tmp_path)
    row_18 = next(f for f in reset_job.findings if f["row_id"] == 18)
    assert row_18["review_status"] == "pending"
    assert row_18["reviewer_note"] == ""
    assert reset_job.human_review_required is True

    reset_summary = Path(reset_job.output_md).read_text(encoding="utf-8")
    assert "Row 18: Severity changed - Production secret in application-prod.yml" in reset_summary
    assert "Approved after confirming credential rotation plan." not in reset_summary

    wb = openpyxl.load_workbook(reset_job.output_xlsx)
    ws = wb["Reviewed Findings"]
    xlsx_row_18 = next(row for row in ws.iter_rows(values_only=True) if row[1] == 18)
    assert xlsx_row_18[10] == "pending"
    assert xlsx_row_18[11] is None
